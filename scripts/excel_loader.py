#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
excel_loader.py —— 从桌面虚拟Excel直接提取数据，输出5张表JSON（驼峰字段）
供 mock_server.js 启动时调用，实现"即启动即用">
架构：读Excel → 列映射 → 表JSON → stdout

输出 stdout: {"ok":true, "tables":{ "pxe_leader_info":[...], "pxe_training_plan":[...], ... }}
字段命名与后端实体一致（驼峰），与 train_api.js / field_map.js 对接：
  planId / className / sessionNo / empId / unitName / quotaCount ...  + 系统字段 delFlag/createTime
"""
import sys, json, re, os
from openpyxl import load_workbook

# 桌面虚拟Excel路径（固化为已知位置）
EXCEL_DIR = r"C:/Users/admin/Desktop/工作/培小e"
FILES = {
    'leader': os.path.join(EXCEL_DIR, '领导人员信息表-山东17地市虚拟.xlsx'),
    'plan': os.path.join(EXCEL_DIR, '2027年度四级领导人员培训计划表（虚拟）.xlsx'),
    'assignment': os.path.join(EXCEL_DIR, '参培人员分配表（2027年度）.xlsx'),
}

now = "2026-10-09 00:00:00"

def cell(v):
    return '' if v is None else v

def session_no(txt):
    m = re.search(r'第([一二三四五六七八九十]+)期', str(txt or ''))
    cn = {'一':1,'二':2,'三':3,'四':4,'五':5,'六':6,'七':7,'八':8,'九':9,'十':10}
    if not m: return None
    s = m.group(1)
    if s == '十': return 10
    if '十' in s:
        return cn.get(s[0],1)*10 + (cn.get(s[-1],0) if s[-1]!='十' else 0)
    return cn.get(s)

def parse_dates(txt):
    txt = str(txt or '')
    m = re.match(r'(\d{4})年(\d{1,2})月(\d{1,2})日', txt)
    if not m: return ('','')
    y=int(m.group(1)); m1=int(m.group(2)); d1=int(m.group(3))
    m2 = re.search(r'(\d{1,2})月(\d{1,2})日$', txt)
    if m2:
        return (f'{y}-{m1:02d}-{d1:02d}', f'{y}-{int(m2.group(1)):02d}-{int(m2.group(2)):02d}')
    return (f'{y}-{m1:02d}-{d1:02d}', '')

def plan_rule(class_name):
    if '新提任四级正职' in class_name:
        return {"cadreLevel":["正科级","副科级"],"target":"2027年新提任四级正职"}, "2027年新提任四级正职领导人员"
    if '新提任四级副职' in class_name:
        return {"cadreLevel":["副科级","四级职员"],"target":"2027年新提任四级副职"}, "2027年新提任四级副职领导人员"
    if '县公司领导' in class_name:
        return {"cadreLevel":["正科级","副科级"]}, "县公司领导班子成员"
    return {"cadreLevel":["正科级","副科级","四级职员"]}, "四级领导人员"

def load_leaders(path):
    """领导人员信息表 → pxe_leader_info（驼峰字段）"""
    if not os.path.exists(path):
        return {"ok": False, "error": f"找不到人员表: {path}"}
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    leaders = []
    for i, r in enumerate(rows[2:]):  # 跳过标题行(0)+表头行(1)
        if not r or all(c is None or str(c).strip()=='' for c in r): continue
        seq = r[0]; unit = cell(r[1]); name = cell(r[2]); eid = str(cell(r[3])).zfill(8)
        gender = cell(r[4]); age = r[5]; assist = cell(r[6]); b_role = cell(r[7])
        level = cell(r[8]); grade = cell(r[9]); party = cell(r[10])
        trained = cell(r[11]); hours = r[12] if r[12] is not None else 0; prod = cell(r[13])
        id_card = '37010' + str(eid).zfill(8) + str(abs(hash(eid))%10) + 'X'
        leaders.append({
            'leaderId': i+1, 'seqNo': seq, 'regionName': '山东省',
            'unitName': unit, 'empName': name, 'empId': eid,
            'idCard': id_card, 'gender': gender, 'age': age,
            'assistDate': assist, 'bRole': b_role, 'cadreLevel': level,
            'staffGrade': grade, 'isPartySec': party,
            'trainedRecords': trained, 'totalHours': hours,
            'isProdHead': prod, 'remark': '虚拟数据-领导人员信息表',
            'delFlag': 0, 'createTime': now,
        })
    return {"ok": True, "rows": leaders, "count": len(leaders)}

def load_plans(path):
    """参培人员分配表(按期次汇总) → pxe_training_plan（驼峰字段）"""
    if not os.path.exists(path):
        return {"ok": False, "error": f"找不到分配表: {path}"}
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb['按期次汇总']
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    plans = []
    for r in rows[2:]:
        if not r or all(c is None or str(c).strip()=='' for c in r): continue
        class_name = cell(r[1]); sess = session_no(r[2])
        dstart, dend = parse_dates(r[3]); pcount = r[4] if r[4] is not None else 0
        if not class_name: continue
        pid = len(plans) + 1
        rjson, target = plan_rule(class_name)
        plans.append({
            'planId': pid, 'batchId': f'2027-{pid:03d}', 'planYear': 2027,
            'seqNo': pid, 'className': class_name, 'sessionNo': sess,
            'trainPeriod': '脱产培训', 'startDate': dstart, 'endDate': dend,
            'trainPlace': '公司党校', 'planCount': pcount,
            'targetGroup': target, 'ruleDesc': json.dumps(rjson, ensure_ascii=False),
            'ruleKey': f'R{pid}', 'matchingRules': rjson,
            'remark': '虚拟数据-按期次汇总', 'delFlag': 0, 'createTime': now,
        })
    return {"ok": True, "rows": plans, "count": len(plans)}

def load_records(path, plans):
    """参培人员分配表(分配总表) → pxe_train_record（驼峰字段）"""
    if not os.path.exists(path):
        return {"ok": False, "error": f"找不到分配表: {path}"}
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb['分配总表']
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    # 班次+期次 → planId 索引
    plan_index = {}
    for p in plans:
        plan_index[(p['className'], p['sessionNo'])] = p['planId']
    records = []
    for i, r in enumerate(rows[2:]):
        if not r or all(c is None or str(c).strip()=='' for c in r): continue
        class_name = cell(r[1]); sess = session_no(r[2])
        dstart, dend = parse_dates(r[3])
        unit = cell(r[4]); name = cell(r[5]); eid = str(cell(r[6])).zfill(8)
        hours = r[10] if r[10] is not None else 0; remark = cell(r[12])
        pid = plan_index.get((class_name, sess))
        if pid is None: continue
        records.append({
            'recordId': i+1, 'planId': pid, 'empId': eid, 'empName': name,
            'planYear': 2027, 'className': class_name, 'classType': '调训班',
            'sessionNo': sess, 'startDate': dstart, 'endDate': dend,
            'hours': hours if hours else 40, 'unitAtTime': unit,
            'dataSource': '虚拟分配', 'remark': remark,
            'delFlag': 0, 'createTime': now,
        })
    return {"ok": True, "rows": records, "count": len(records)}

def load_quotas(plans):
    """按单位生成配额模拟数据 → pxe_quota_alloc（驼峰字段）"""
    units17 = ['济南公司','青岛公司','淄博公司','枣庄公司','东营公司',
               '烟台公司','潍坊公司','济宁公司','泰安公司','威海公司',
               '日照公司','莱芜公司','临沂公司','德州公司','聊城公司',
               '滨州公司','菏泽公司']
    quotas = []
    qid = 0
    for p in plans:
        base = max(p['planCount'] // len(units17), 1) if p['planCount'] else 2
        for u in units17:
            qid += 1
            quotas.append({
                'quotaId': qid, 'planId': p['planId'],
                'unitName': u, 'quotaCount': base,
                'remark': '虚拟数据-配额模拟', 'delFlag': 0, 'createTime': now,
            })
    return {"ok": True, "rows": quotas, "count": len(quotas)}

def load_rules():
    """通用规则表 → pxe_training_rule（驼峰字段）
    编码与 ai_advisor.js 的 RULE_EXECUTORS 一一对应，可直接按 ruleCode 命中执行器。"""
    rules = [
        {'ruleType': '硬规则', 'ruleCode': 'HARD_AB_ROLE', 'ruleName': '硬规则一：AB角不得同出',
         'ruleContent': '互为AB角的两名领导人员，同期次培训时间重叠时不得同时参培', 'priority': 1, 'enabled': '1'},
        {'ruleType': '硬规则', 'ruleCode': 'HARD_REPEAT_TRAIN', 'ruleName': '硬规则二：重复参培剔除',
         'ruleContent': '2027年度原则上只安排参加1个脱产培训班次，2026年度已参加同类别班次的不再安排', 'priority': 2, 'enabled': '1'},
        {'ruleType': '硬规则', 'ruleCode': 'HARD_PROD_HEAD_SEASON', 'ruleName': '硬规则三：生产分管负责人避开春秋检',
         'ruleContent': '生产分管负责人的培训期次需避开春检（3-6月）与秋检（9-11月）', 'priority': 3, 'enabled': '1'},
        {'ruleType': '硬规则', 'ruleCode': 'HARD_TARGET_LIMIT', 'ruleName': '硬规则四：参培对象限制',
         'ruleContent': '距退休5年内（55岁以上）不安排进修班；新提任班仅限对应级别人员', 'priority': 4, 'enabled': '1'},
        {'ruleType': '软规则', 'ruleCode': 'SOFT_RECENT_UNTRAINED', 'ruleName': '优先规则一：近3年未参培优先',
         'ruleContent': '近3年未参加任何培训的人员优先安排', 'priority': 6, 'enabled': '1'},
        {'ruleType': '软规则', 'ruleCode': 'SOFT_HOURS_GAP', 'ruleName': '优先规则二：学时缺口优先',
         'ruleContent': '累计学时偏低（低于100）的人员优先安排', 'priority': 7, 'enabled': '1'},
        {'ruleType': '软规则', 'ruleCode': 'SOFT_PARTY_SEC_ONLY', 'ruleName': '优先规则三：党支部书记专场',
         'ruleContent': '党支部书记统一安排在进修班第五期专场', 'priority': 8, 'enabled': '1'},
    ]
    rows = []
    for i, r in enumerate(rules):
        rows.append(dict(r, ruleId=i+1, priority=r['priority'], remark='虚拟数据-规则表',
                         delFlag=0, createTime=now))
    return {"ok": True, "rows": rows, "count": len(rows)}

def main():
    # 检查Excel目录
    if not os.path.exists(EXCEL_DIR):
        print(json.dumps({"ok": False, "error": f"桌面Excel目录不存在: {EXCEL_DIR}"}))
        sys.exit(1)

    # 检查必需文件存在
    for key, fpath in FILES.items():
        if not os.path.exists(fpath):
            print(json.dumps({"ok": False, "error": f"缺少文件: {fpath}"}))
            sys.exit(1)

    # 1. 加载人员信息
    result = load_leaders(FILES['leader'])
    if not result['ok']:
        print(json.dumps(result)); sys.exit(1)
    leaders = result['rows']

    # 2. 加载培训计划（按期次汇总）
    result = load_plans(FILES['assignment'])
    if not result['ok']:
        print(json.dumps(result)); sys.exit(1)
    plans = result['rows']

    # 3. 加载培训记录（分配总表）
    result = load_records(FILES['assignment'], plans)
    if not result['ok']:
        print(json.dumps(result)); sys.exit(1)
    records = result['rows']

    # 4. 生成单位配额
    result = load_quotas(plans)
    if not result['ok']:
        print(json.dumps(result)); sys.exit(1)
    quotas = result['rows']

    # 5. 生成通用规则
    result = load_rules()
    if not result['ok']:
        print(json.dumps(result)); sys.exit(1)
    rules = result['rows']

    output = {
        "ok": True,
        "tables": {
            "pxe_leader_info": leaders,
            "pxe_training_plan": plans,
            "pxe_train_record": records,
            "pxe_quota_alloc": quotas,
            "pxe_training_rule": rules,
        },
        "counts": {
            "pxe_leader_info": len(leaders),
            "pxe_training_plan": len(plans),
            "pxe_train_record": len(records),
            "pxe_quota_alloc": len(quotas),
            "pxe_training_rule": len(rules),
        }
    }
    print(json.dumps(output, ensure_ascii=False, default=str))

if __name__ == '__main__':
    main()
